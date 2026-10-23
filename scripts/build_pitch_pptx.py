"""SchemeNavigator pitch — exact user content, light theme, 7 slides."""
from __future__ import annotations

import os
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BG_IMG = os.path.abspath(os.path.join(os.path.dirname(__file__), "assets", "website_bg_grid.png"))

NAVY = RGBColor(0x0F, 0x2C, 0x4C)
TEAL = RGBColor(0x0D, 0x94, 0x88)
TEAL_DK = RGBColor(0x0B, 0x76, 0x6C)
INK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x47, 0x55, 0x69)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF7, 0xF9, 0xFC)
LINE = RGBColor(0xE2, 0xE8, 0xF0)
GOLD = RGBColor(0xC2, 0x41, 0x0C)
AMBER = RGBColor(0xD9, 0xD7, 0x06)
SKY = RGBColor(0x03, 0x69, 0xA1)
VIOLET = RGBColor(0x6D, 0x28, 0xD9)
ROSE = RGBColor(0xBE, 0x12, 0x3C)
SLATE = RGBColor(0xF1, 0xF5, 0xF9)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)

W, H = Inches(13.333), Inches(7.5)


def set_run(run, size, bold=False, color=INK, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = "Calibri"


def write_box(shape, lines, valign=MSO_ANCHOR.TOP):
    tf = shape.text_frame
    tf.word_wrap = True
    try:
        tf.vertical_anchor = valign
    except Exception:
        pass
    for i, item in enumerate(lines):
        text, size, bold, color, italic, align, space_after = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = text
        set_run(run, size, bold, color, italic)


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def fill_line(shape, fill_c, line_c=LINE, pt=1.15):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_c
    shape.line.color.rgb = line_c
    shape.line.width = Pt(pt)


def shadow(shape):
    spPr = shape._element.spPr
    for child in list(spPr):
        if "effectLst" in child.tag:
            spPr.remove(child)
    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    el = etree.SubElement(spPr, "{%s}effectLst" % ns)
    outer = etree.SubElement(
        el,
        "{%s}outerShdw" % ns,
        {
            "blurRad": "50800",
            "dist": "38100",
            "dir": "16200000",
            "algn": "tl",
            "rotWithShape": "0",
        },
    )
    srgb = etree.SubElement(outer, "{%s}srgbClr" % ns, {"val": "0F2C4C"})
    etree.SubElement(srgb, "{%s}alpha" % ns, {"val": "16000"})


def rect(slide, l, t, w, h, color=None, rounded=True, line=None, sh=False, adj=0.08):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, l, t, w, h)
    if color is None:
        s.fill.background()
        s.line.fill.background()
    elif line:
        fill_line(s, color, line)
    else:
        fill(s, color)
    if rounded:
        try:
            s.adjustments[0] = adj
        except Exception:
            pass
    if sh:
        shadow(s)
    return s


def pill(slide, l, t, w, h, text, bg, fg=WHITE, size=11):
    s = rect(slide, l, t, w, h, bg, adj=0.5)
    write_box(s, [(text, size, True, fg, False, PP_ALIGN.CENTER, 0)], MSO_ANCHOR.MIDDLE)
    return s


def footer(slide, n):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.28), W, Inches(0.22))
    fill(bar, NAVY)
    tf = bar.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = "  SchemeNavigator  ·  Discover. Understand. Apply.                                              %d / 7" % n
    set_run(r, 9, False, WHITE)


def base(prs, accent):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if os.path.exists(BG_IMG):
        s.shapes.add_picture(BG_IMG, 0, 0, W, H)
    else:
        fill(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H), CREAM)
    fill(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), H), accent)
    fill(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.08)), accent)
    return s


prs = Presentation()
prs.slide_width = W
prs.slide_height = H


# ===================== SLIDE 1 =====================
s = base(prs, TEAL)
write_box(
    rect(s, Inches(0.55), Inches(1.15), Inches(12.2), Inches(0.95), None),
    [("SchemeNavigator", 48, True, NAVY, False, PP_ALIGN.CENTER, 0)],
    MSO_ANCHOR.MIDDLE,
)
write_box(
    rect(s, Inches(0.55), Inches(2.1), Inches(12.2), Inches(0.55), None),
    [("Discover. Understand. Apply.", 24, True, TEAL, False, PP_ALIGN.CENTER, 0)],
    MSO_ANCHOR.MIDDLE,
)
write_box(
    rect(s, Inches(0.55), Inches(2.65), Inches(12.2), Inches(0.45), None),
    [("AI-powered navigation for Government Schemes", 18, False, MUTED, False, PP_ALIGN.CENTER, 0)],
    MSO_ANCHOR.MIDDLE,
)

team = rect(s, Inches(1.8), Inches(3.4), Inches(9.7), Inches(1.85), WHITE, line=LINE, sh=True)
write_box(
    team,
    [
        ("Team [Name]", 20, True, NAVY, False, PP_ALIGN.CENTER, 10),
        ("[Member 1]  •  [Member 2]  •  [Member 3]  •  [Member 4]", 16, False, INK, False, PP_ALIGN.CENTER, 0),
    ],
    MSO_ANCHOR.MIDDLE,
)

banner = rect(s, Inches(1.8), Inches(5.55), Inches(9.7), Inches(0.7), NAVY)
write_box(
    banner,
    [("Citizen-first  •  Privacy-first  •  Explainable", 16, True, WHITE, True, PP_ALIGN.CENTER, 0)],
    MSO_ANCHOR.MIDDLE,
)
footer(s, 1)


# ===================== SLIDE 2 =====================
s = base(prs, ROSE)
write_box(
    rect(s, Inches(0.5), Inches(0.28), Inches(12.3), Inches(0.95), None),
    [("India has schemes. Citizens lack navigation.", 28, True, NAVY, False, PP_ALIGN.LEFT, 0)],
    MSO_ANCHOR.MIDDLE,
)

blocks = [
    ("DISCOVER", "Don't know which schemes exist.", TEAL),
    ("UNDERSTAND", "Eligibility & documents are difficult to decode.", AMBER),
    ("TRUST", "Hard to identify the genuine application portal.", GOLD),
]
for i, (lab, body, col) in enumerate(blocks):
    x = Inches(0.5 + i * 4.2)
    c = rect(s, x, Inches(1.45), Inches(3.95), Inches(3.15), WHITE, line=LINE, sh=True)
    fill(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.45), Inches(0.12), Inches(3.15)), col)
    write_box(
        c,
        [
            (lab, 16, True, col, False, PP_ALIGN.LEFT, 12),
            (body, 18, True, NAVY, False, PP_ALIGN.LEFT, 0),
        ],
    )

res = rect(s, Inches(0.5), Inches(4.9), Inches(12.3), Inches(1.85), WHITE, line=LINE, sh=True)
write_box(
    res,
    [
        ("The Result", 14, True, ROSE, False, PP_ALIGN.CENTER, 10),
        ("Eligible citizens can miss benefits that already exist.", 22, True, NAVY, False, PP_ALIGN.CENTER, 0),
    ],
    MSO_ANCHOR.MIDDLE,
)
footer(s, 2)


# ===================== SLIDE 3 =====================
s = base(prs, TEAL)
write_box(
    rect(s, Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.75), None),
    [("From Citizen Profile → Right Scheme → Right Action", 24, True, NAVY, False, PP_ALIGN.LEFT, 0)],
    MSO_ANCHOR.MIDDLE,
)

steps = [
    ("1. PROFILE", "Tell us about yourself", TEAL),
    ("2. MATCH", "Rule-based scheme matching", SKY),
    ("3. EXPLAIN", "Why you match + benefits + documents", VIOLET),
    ("4. ASSIST", "AI Advisor • Voice • 12+ Languages", AMBER),
    ("5. ACT", "Official Portal • Tracker • Kendra", GOLD),
]
for i, (lab, body, col) in enumerate(steps):
    x = Inches(0.38 + i * 2.58)
    c = rect(s, x, Inches(1.25), Inches(2.42), Inches(3.85), WHITE, line=LINE, sh=True)
    fill(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.25), Inches(2.42), Inches(0.1)), col)
    write_box(
        c,
        [
            (lab, 14, True, col, False, PP_ALIGN.CENTER, 10),
            (body, 14, True, NAVY, False, PP_ALIGN.CENTER, 0),
        ],
    )
    if i < 4:
        ar = s.shapes.add_shape(
            MSO_SHAPE.RIGHT_ARROW,
            Inches(0.38 + i * 2.58 + 2.44),
            Inches(3.05),
            Inches(0.14),
            Inches(0.22),
        )
        fill(ar, TEAL)

banner = rect(s, Inches(0.5), Inches(5.4), Inches(12.3), Inches(1.35), NAVY)
write_box(
    banner,
    [("Rules decide. AI explains.", 28, True, WHITE, False, PP_ALIGN.CENTER, 0)],
    MSO_ANCHOR.MIDDLE,
)
footer(s, 3)


# ===================== SLIDE 4 =====================
s = base(prs, BLUE)
write_box(
    rect(s, Inches(0.5), Inches(0.28), Inches(12.3), Inches(0.55), None),
    [("IBM BoB × SchemeNavigator", 14, True, BLUE, False, PP_ALIGN.LEFT, 0)],
)
write_box(
    rect(s, Inches(0.5), Inches(0.72), Inches(12.3), Inches(0.7), None),
    [("AI-Assisted Development. Human-Controlled Decisions.", 24, True, NAVY, False, PP_ALIGN.LEFT, 0)],
    MSO_ANCHOR.MIDDLE,
)

pipe = [
    ("UNDERSTAND", "Codebase & requirements"),
    ("PLAN", "Break features into tasks"),
    ("BUILD", "Agent-assisted implementation"),
    ("DEBUG & ITERATE", "Test, refine & improve"),
]
for i, (lab, body) in enumerate(pipe):
    y = Inches(1.55 + i * 1.05)
    c = rect(s, Inches(2.15), y, Inches(8.95), Inches(0.88), WHITE, line=LINE, sh=True)
    write_box(
        c,
        [
            (lab, 15, True, TEAL_DK, False, PP_ALIGN.CENTER, 4),
            (body, 14, False, MUTED, False, PP_ALIGN.CENTER, 0),
        ],
        MSO_ANCHOR.MIDDLE,
    )
    if i < 3:
        dn = s.shapes.add_shape(
            MSO_SHAPE.DOWN_ARROW,
            Inches(6.45),
            Inches(1.55 + i * 1.05 + 0.88),
            Inches(0.22),
            Inches(0.16),
        )
        fill(dn, TEAL)

bot = rect(s, Inches(0.5), Inches(5.75), Inches(12.3), Inches(1.0), NAVY)
write_box(
    bot,
    [
        ("IBM BoB helped us move from idea → working product.", 16, True, WHITE, False, PP_ALIGN.CENTER, 6),
        ("Human review & validation at every step", 14, False, RGBColor(0xCB, 0xD5, 0xE1), False, PP_ALIGN.CENTER, 0),
    ],
    MSO_ANCHOR.MIDDLE,
)
footer(s, 4)


# ===================== SLIDE 5 =====================
s = base(prs, VIOLET)
write_box(
    rect(s, Inches(0.5), Inches(0.28), Inches(12.3), Inches(0.75), None),
    [("Not another scheme directory.", 28, True, NAVY, False, PP_ALIGN.LEFT, 0)],
    MSO_ANCHOR.MIDDLE,
)

diffs = [
    ("EXPLAINABLE", "Why you match — not just what you match", TEAL),
    ("MULTILINGUAL", "12+ languages • Voice interaction", SKY),
    ("PRIVACY-FIRST", "No Aadhaar • No PAN • No OTP", GOLD),
    ("OFFICIAL-FIRST", "Apply through official government portals", VIOLET),
]
for i, (lab, body, col) in enumerate(diffs):
    x = Inches(0.42 + i * 3.2)
    c = rect(s, x, Inches(1.25), Inches(3.05), Inches(3.45), WHITE, line=LINE, sh=True)
    fill(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.25), Inches(3.05), Inches(0.1)), col)
    write_box(
        c,
        [
            (lab, 14, True, col, False, PP_ALIGN.LEFT, 12),
            (body, 16, True, NAVY, False, PP_ALIGN.LEFT, 0),
        ],
    )

bot = rect(s, Inches(0.42), Inches(5.05), Inches(12.45), Inches(1.7), WHITE, line=LINE, sh=True)
write_box(
    bot,
    [("Digital Navigation + Last-Mile Support", 24, True, NAVY, False, PP_ALIGN.CENTER, 0)],
    MSO_ANCHOR.MIDDLE,
)
footer(s, 5)


# ===================== SLIDE 6 =====================
s = base(prs, TEAL)
write_box(
    rect(s, Inches(0.5), Inches(0.28), Inches(12.3), Inches(0.7), None),
    [("A Complete Citizen Journey", 28, True, NAVY, False, PP_ALIGN.LEFT, 0)],
    MSO_ANCHOR.MIDDLE,
)

stats = [
    ("3,866+", "Schemes"),
    ("12+", "Languages"),
    ("0–100", "Explainable Matching"),
    ("AI", "Advisor"),
    ("Voice", "Interaction"),
    ("Kendra", "Finder"),
]
for i, (a, b) in enumerate(stats):
    x = Inches(0.42 + i * 2.13)
    c = rect(s, x, Inches(1.25), Inches(2.0), Inches(2.35), WHITE, line=LINE, sh=True)
    write_box(
        c,
        [
            (a, 22, True, TEAL, False, PP_ALIGN.CENTER, 8),
            (b, 13, True, MUTED, False, PP_ALIGN.CENTER, 0),
        ],
        MSO_ANCHOR.MIDDLE,
    )

j = rect(s, Inches(0.42), Inches(3.95), Inches(12.45), Inches(2.55), NAVY)
write_box(
    j,
    [
        ("Discover → Match → Understand → Apply → Track", 22, True, WHITE, False, PP_ALIGN.CENTER, 12),
        ("LIVE PRODUCT DEMO", 16, True, RGBColor(0x5E, 0xEA, 0xD4), False, PP_ALIGN.CENTER, 0),
    ],
    MSO_ANCHOR.MIDDLE,
)
footer(s, 6)


# ===================== SLIDE 7 =====================
s = base(prs, SKY)
write_box(
    rect(s, Inches(0.5), Inches(0.22), Inches(12.3), Inches(0.7), None),
    [("Where SchemeNavigator Goes Next", 26, True, NAVY, False, PP_ALIGN.LEFT, 0)],
    MSO_ANCHOR.MIDDLE,
)

pillars = [
    ("LIVE SCHEME INTELLIGENCE", "Automated updates from authoritative sources", TEAL),
    ("GOVERNMENT OPPORTUNITIES & EVENTS", "Hackathons • Competitions • Fellowships • Government Events", SKY),
    ("LAST-MILE SCALE", "CSCs • NGOs • Colleges • State Missions", VIOLET),
    ("SUSTAINABLE REVENUE", "B2B / White-label • API • CSR • Institutional Support\nCitizen access remains free", GOLD),
]
for i, (lab, body, col) in enumerate(pillars):
    x = Inches(0.42 + (i % 2) * 6.4)
    y = Inches(1.05 + (i // 2) * 2.25)
    c = rect(s, x, y, Inches(6.15), Inches(2.1), WHITE, line=LINE, sh=True)
    fill(s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.1), Inches(2.1)), col)
    write_box(
        c,
        [
            (lab, 13, True, col, False, PP_ALIGN.LEFT, 8),
            (body, 16, True, NAVY, False, PP_ALIGN.LEFT, 0),
        ],
    )

bot = rect(s, Inches(0.42), Inches(5.65), Inches(12.45), Inches(1.1), NAVY)
write_box(
    bot,
    [("From Scheme Discovery → Government Opportunity Discovery", 18, True, WHITE, False, PP_ALIGN.CENTER, 0)],
    MSO_ANCHOR.MIDDLE,
)
footer(s, 7)

out = r"C:\Users\divya\Downloads\Scheme-Navigator-App-master\Scheme-Navigator-App-master\SchemeNavigator_15min_Pitch.pptx"
try:
    prs.save(out)
    print("Wrote", out)
except PermissionError:
    alt = r"C:\Users\divya\Downloads\Scheme-Navigator-App-master\Scheme-Navigator-App-master\SchemeNavigator_Pitch_Deck.pptx"
    prs.save(alt)
    print("Locked by PowerPoint! Wrote to alternate:", alt)
